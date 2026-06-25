# app/services/google_drive/drive_service.py

import io
import logging

import requests

logger = logging.getLogger(__name__)

_PDF_MIME = "application/pdf"
_DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
_XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

_SKIP_MIMES = {
    "application/zip",
    "application/x-zip-compressed",
    "application/epub+zip",
    "application/octet-stream",
}


def _strip_nul(text: str) -> str:
    """Remove NUL bytes — PostgreSQL rejects strings containing 0x00."""
    return text.replace("\x00", "")


def _extract_pdf_text(content: bytes) -> str:
    """Extract plain text from PDF bytes using pypdf."""
    try:
        import pypdf
        reader = pypdf.PdfReader(io.BytesIO(content))
        pages = []
        for page in reader.pages:
            try:
                pages.append(page.extract_text() or "")
            except Exception:
                pass
        return "\n".join(pages)
    except ImportError:
        logger.warning("[DRIVE] pypdf not installed — lossy decode fallback for PDF")
        return content.decode("utf-8", errors="ignore")
    except Exception as e:
        logger.warning(f"[DRIVE] PDF extraction failed: {e}")
        return content.decode("utf-8", errors="ignore")


def _extract_docx_text(content: bytes) -> str:
    """Extract plain text from DOCX bytes using python-docx."""
    try:
        import docx
        doc = docx.Document(io.BytesIO(content))
        return "\n".join(p.text for p in doc.paragraphs if p.text.strip())
    except ImportError:
        logger.warning("[DRIVE] python-docx not installed — lossy decode fallback for DOCX")
        return content.decode("utf-8", errors="ignore")
    except Exception as e:
        logger.warning(f"[DRIVE] DOCX extraction failed: {e}")
        return content.decode("utf-8", errors="ignore")


def _extract_xlsx_text(content: bytes) -> str:
    """Extract cell text from XLSX using openpyxl."""
    try:
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        rows = []
        for sheet in wb.worksheets:
            rows.append(f"[Sheet: {sheet.title}]")
            for row in sheet.iter_rows(values_only=True):
                cell_texts = [str(c) for c in row if c is not None and str(c).strip()]
                if cell_texts:
                    rows.append("\t".join(cell_texts))
        return "\n".join(rows)
    except ImportError:
        logger.warning("[DRIVE] openpyxl not installed — skipping XLSX")
        return ""
    except Exception as e:
        logger.warning(f"[DRIVE] XLSX extraction failed: {e}")
        return ""


def _extract_text(content: bytes, mime_type: str, file_name: str) -> str:
    """
    Route binary content to the right extractor based on MIME type.
    Always returns a clean string — never raises.
    """
    if not content:
        return ""

    name_lower = file_name.lower()

    # Skip binary containers with no useful plain text (zip, epub)
    if mime_type in _SKIP_MIMES or name_lower.endswith((".zip", ".epub")):
        logger.info(f"[DRIVE] Skipping binary container: {file_name}")
        return ""

    if mime_type == _XLSX_MIME or name_lower.endswith((".xlsx", ".xls")):
        return _strip_nul(_extract_xlsx_text(content))

    if mime_type == _PDF_MIME or name_lower.endswith(".pdf"):
        return _strip_nul(_extract_pdf_text(content))

    if mime_type == _DOCX_MIME or name_lower.endswith(".docx"):
        return _strip_nul(_extract_docx_text(content))

    # Plain text (txt, csv, json, etc.) — decode as UTF-8, drop invalid bytes
    return _strip_nul(content.decode("utf-8", errors="ignore"))


class DriveService:

    @staticmethod
    def get_files(access_token: str):
        response = requests.get(
            "https://www.googleapis.com/drive/v3/files",
            params={
                "pageSize": 100,
                "fields": (
                    "files("
                    "id,"
                    "name,"
                    "mimeType,"
                    "owners,"
                    "size,"
                    "modifiedTime,"
                    "webViewLink"
                    ")"
                )
            },
            headers={"Authorization": f"Bearer {access_token}"}
        )
        response.raise_for_status()
        return response.json()

    @staticmethod
    def download_file(
        access_token: str,
        file_id: str,
        mime_type: str = "",
        file_name: str = "",
    ) -> str:
        """
        Download a file from Google Drive and return its text content.
        Handles PDF, DOCX, XLSX, TXT, CSV correctly.
        ZIP/EPUB return "". Never raises UnicodeDecodeError.
        """
        response = requests.get(
            f"https://www.googleapis.com/drive/v3/files/{file_id}?alt=media",
            headers={"Authorization": f"Bearer {access_token}"}
        )
        response.raise_for_status()
        return _extract_text(response.content, mime_type, file_name)

    @staticmethod
    def export_google_doc(access_token: str, file_id: str) -> str:
        """Export a Google Doc as plain text — unchanged."""
        response = requests.get(
            f"https://www.googleapis.com/drive/v3/files/{file_id}/export",
            params={"mimeType": "text/plain"},
            headers={"Authorization": f"Bearer {access_token}"}
        )
        response.raise_for_status()
        return response.content.decode("utf-8-sig")